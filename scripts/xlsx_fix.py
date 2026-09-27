# -*- coding: utf-8 -*-
"""卖家精灵导出 xlsx 的 openpyxl 兼容修复

问题: 卖家精灵生成的 xl/drawings/drawingN.xml 里 anchor 用了非标字段
      editAs="undefined", openpyxl 严格按 OOXML 校验会抛:
        ValueError: Value must be one of {'absolute','twoCell','oneCell'}

用法:
  from xlsx_fix import load_workbook_safe
  wb = load_workbook_safe('Search(xxx)-162-US-20260910.xlsx')   # 可写模式

  # 命令行自检
  python xlsx_fix.py <file.xlsx>
"""
import os
import re
import zipfile


def fix_edit_as(xlsx_path):
    """解压 -> 把 editAs="undefined" 换成 "oneCell" -> 重打包，返回临时文件路径。"""
    src = xlsx_path
    tmp = src.replace('.xlsx', '.fix%s.xlsx' % os.getpid())
    with zipfile.ZipFile(src, 'r') as zin, \
            zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith('xl/drawings/') and item.filename.endswith('.xml'):
                txt = data.decode('utf-8', errors='ignore')
                txt = re.sub(r'editAs="undefined"', 'editAs="oneCell"', txt)
                data = txt.encode('utf-8')
            zout.writestr(item, data)
    return tmp


def load_workbook_safe(xlsx_path, **kw):
    import openpyxl
    try:
        return openpyxl.load_workbook(xlsx_path, **kw)
    except ValueError:
        tmp = fix_edit_as(xlsx_path)
        try:
            return openpyxl.load_workbook(tmp, **kw)
        finally:
            try:
                os.unlink(tmp)
            except OSError:
                pass


def clear_inherited_images(ws):
    """重写文件时先清掉继承的旧图，否则新旧图锚在同一区域会重叠。"""
    try:
        ws._images.clear()
    except AttributeError:
        ws._images = []


if __name__ == '__main__':
    import sys
    p = sys.argv[1]
    wb = load_workbook_safe(p)
    ws = wb.active
    print('OK', p, '| sheets=', wb.sheetnames, '| dims=', ws.max_row, 'x', ws.max_column)
